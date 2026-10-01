"""
BFSI Comps & Quarterly Results Tracker
--------------------------------------
1. Ranks listed private banks on valuation (P/E, P/B) and profit growth, with medians by size bucket.
2. Fits P/E against profit growth (OLS, outliers removed by the IQR rule) to flag banks that look
   cheap or expensive for their growth.
3. Fits P/B against ROE - the standard bank valuation lens - once at least 5 banks have book value
   and ROE filled in data/peers.csv.
4. Tracks Karur Vysya Bank's quarterly results (YoY / QoQ) and builds a results-day sheet that
   compares reported numbers with model-implied estimates.

Run:  python bfsi_toolkit.py            (reads ./data, writes ./output)
Needs: pandas, numpy, matplotlib, openpyxl
"""
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as col

FOCUS = "Karur Vysya Bank"
LARGE_CAP, MID_CAP = 100_000, 35_000          # Rs crore market-cap cut-offs
DATA, OUT = "data", "output"
NAVY, KVB_COL, GREY = "#1F3864", "#C00000", "#8C8C8C"


def size_bucket(mcap):
    return "Large" if mcap >= LARGE_CAP else ("Mid" if mcap >= MID_CAP else "Small")


def iqr_ok(s, k=1.5):
    q1, q3 = s.quantile([0.25, 0.75])
    return s.between(q1 - k * (q3 - q1), q3 + k * (q3 - q1))


def ols(x, y):
    b, a = np.polyfit(x, y, 1)
    pred = a + b * x
    r2 = 1 - ((y - pred) ** 2).sum() / ((y - y.mean()) ** 2).sum()
    return a, b, r2


def build_comps(peers):
    df = peers.copy()
    df["size"] = df["mcap_cr"].map(size_bucket)
    df["pb"] = df["cmp"] / df["bvps"]
    df["earnings_yield_pct"] = 100 / df["pe_ttm"]
    df["pe_pctile"] = df["pe_ttm"].rank(pct=True) * 100          # 100 = most expensive
    df["growth_pctile"] = df["pat_yoy_pct"].rank(pct=True) * 100  # 100 = fastest growth
    df["in_fit"] = iqr_ok(df["pe_ttm"]) & iqr_ok(df["pat_yoy_pct"])
    a, b, r2 = ols(df.loc[df.in_fit, "pat_yoy_pct"].values, df.loc[df.in_fit, "pe_ttm"].values)
    df["pe_fitted"] = a + b * df["pat_yoy_pct"]
    df["pe_residual"] = df["pe_ttm"] - df["pe_fitted"]            # negative = cheap for its growth
    return df.sort_values("mcap_cr", ascending=False), {"a": a, "b": b, "r2": r2, "n": int(df.in_fit.sum())}


def chart_pe_growth(df, fit, path):
    fig, ax = plt.subplots(figsize=(8, 5))
    inn, out = df[df.in_fit], df[~df.in_fit]
    ax.scatter(inn.pat_yoy_pct, inn.pe_ttm, s=40, color=NAVY, label="Used in fit")
    ax.scatter(out.pat_yoy_pct, out.pe_ttm, s=40, facecolors="none", edgecolors=GREY, label="Outlier (excluded)")
    xs = np.linspace(inn.pat_yoy_pct.min(), inn.pat_yoy_pct.max(), 50)
    ax.plot(xs, fit["a"] + fit["b"] * xs, color=GREY, ls="--", lw=1, label=f"OLS fit, R² = {fit['r2']:.2f}")
    for _, r in df.iterrows():
        focus = r.bank == FOCUS
        ax.annotate(r.bank.replace(" Bank", ""), (r.pat_yoy_pct, r.pe_ttm), fontsize=7 if not focus else 8,
                    color=KVB_COL if focus else "#333333", fontweight="bold" if focus else "normal",
                    xytext=(4, 3), textcoords="offset points")
    k = df[df.bank == FOCUS]
    ax.scatter(k.pat_yoy_pct, k.pe_ttm, s=90, color=KVB_COL, zorder=5)
    ax.set_xlabel("Q1 FY27 net profit growth, YoY (%)"); ax.set_ylabel("P/E (TTM, x)")
    ax.set_title("Private banks: P/E vs profit growth", loc="left", fontsize=11, fontweight="bold")
    ax.grid(alpha=0.25); ax.legend(fontsize=8, frameon=False); fig.tight_layout(); fig.savefig(path, dpi=200); plt.close(fig)


def pb_roe(df, path, min_banks=5):
    have = df.dropna(subset=["bvps", "roe_pct"])
    if len(have) < min_banks:
        return None, len(have)
    a, b, r2 = ols(have.roe_pct.values, have.pb.values)
    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.scatter(have.roe_pct, have.pb, color=NAVY)
    xs = np.linspace(have.roe_pct.min(), have.roe_pct.max(), 50); ax.plot(xs, a + b * xs, ls="--", color=GREY)
    for _, r in have.iterrows():
        ax.annotate(r.bank.replace(" Bank", ""), (r.roe_pct, r.pb), fontsize=7, xytext=(4, 3), textcoords="offset points")
    ax.set_xlabel("ROE (%)"); ax.set_ylabel("P/B (x)"); ax.set_title(f"P/B vs ROE (R² = {r2:.2f})", loc="left")
    fig.tight_layout(); fig.savefig(path, dpi=200); plt.close(fig)
    return {"a": a, "b": b, "r2": r2, "n": len(have)}, len(have)


def build_tracker(q):
    q = q.copy()
    for m in ["nii", "other_income", "pbt", "net_profit"]:
        q[f"{m}_yoy"] = q[m] / q[m].shift(4) - 1
        q[f"{m}_qoq"] = q[m] / q[m].shift(1) - 1
    return q


def chart_tracker(q, path):
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(8, 6), sharex=True, gridspec_kw={"height_ratios": [2, 1]})
    x = np.arange(len(q)); w = 0.4
    a1.bar(x - w / 2, q.nii, w, color=NAVY, label="NII (Screener basis)")
    a1.bar(x + w / 2, q.net_profit, w, color=KVB_COL, label="Net profit")
    a1.set_ylabel("Rs crore"); a1.legend(fontsize=8, frameon=False); a1.grid(axis="y", alpha=0.25)
    a1.set_title("Karur Vysya Bank: quarterly NII, profit and asset quality", loc="left", fontsize=11, fontweight="bold")
    a2.plot(x, q.gnpa_pct, marker="o", color=NAVY, label="Gross NPA %"); a2.plot(x, q.nnpa_pct, marker="o", color=GREY, label="Net NPA %")
    a2.set_xticks(x); a2.set_xticklabels(q.quarter, rotation=45, fontsize=8); a2.legend(fontsize=8, frameon=False); a2.grid(alpha=0.25)
    fig.tight_layout(); fig.savefig(path, dpi=200); plt.close(fig)


def write_excel(df, fit, q, est, path):
    wb = Workbook(); H = Font(name="Arial", size=9, bold=True, color="FFFFFF"); N = Font(name="Arial", size=9)
    BLUE = Font(name="Arial", size=9, color="0000FF"); FILL = PatternFill("solid", fgColor="1F3864"); YEL = PatternFill("solid", fgColor="FFFF00")

    def table(ws, frame, fmts, r0=1):
        for j, c in enumerate(frame.columns, 1):
            cell = ws.cell(r0, j, c); cell.font = H; cell.fill = FILL; cell.alignment = Alignment(wrap_text=True, horizontal="center")
            ws.column_dimensions[col(j)].width = 24 if j == 1 else 12
        for i, row in enumerate(frame.itertuples(index=False), r0 + 1):
            for j, v in enumerate(row, 1):
                if isinstance(v, float) and np.isnan(v): v = None
                cell = ws.cell(i, j, v); cell.font = N; cell.number_format = fmts.get(frame.columns[j - 1], "General")

    ws = wb.active; ws.title = "Read_me"
    for i, t in enumerate(["BFSI comps & KVB results tracker - generated by bfsi_toolkit.py",
                           "Comps / PE_vs_Growth: Python output (re-run the script after editing data/peers.csv).",
                           "KVB_Quarterly: raw numbers in blue; growth rows are live formulas.",
                           "Q2FY27_Results_Day: type reported numbers into the yellow cells on 15-Oct-2026; variance and verdict update.",
                           "Data: Screener.in industry table and company pages (30-Sep-2026). Student project - not investment advice."], 1):
        ws.cell(i, 1, t).font = Font(name="Arial", size=10, bold=(i == 1))
    ws.column_dimensions["A"].width = 110

    cols = ["bank", "size", "cmp", "mcap_cr", "pe_ttm", "pe_pctile", "pat_yoy_pct", "growth_pctile", "earnings_yield_pct", "div_yield_pct", "bvps", "roe_pct", "pb"]
    table(wb.create_sheet("Comps"), df[cols], {"cmp": "#,##0.0", "mcap_cr": "#,##0", "pe_ttm": "0.0", "pe_pctile": "0", "pat_yoy_pct": "0.0",
                                              "growth_pctile": "0", "earnings_yield_pct": "0.0", "div_yield_pct": "0.00", "pb": "0.00"})
    w2 = wb.create_sheet("PE_vs_Growth")
    table(w2, df[["bank", "pat_yoy_pct", "pe_ttm", "pe_fitted", "pe_residual", "in_fit"]].sort_values("pe_residual"),
          {"pat_yoy_pct": "0.0", "pe_ttm": "0.0", "pe_fitted": "0.0", "pe_residual": "0.0"})
    w2.cell(23, 1, f"Fit on {fit['n']} banks: P/E = {fit['a']:.2f} + {fit['b']:.3f} x growth(%), R-squared {fit['r2']:.2f}. Negative residual = cheap for its growth.").font = N

    w3 = wb.create_sheet("KVB_Quarterly"); raw = ["nii", "other_income", "pbt", "net_profit", "gnpa_pct", "nnpa_pct"]
    w3.cell(1, 1, "Rs crore").font = H; w3.cell(1, 1).fill = FILL; w3.column_dimensions["A"].width = 26
    for j, qq in enumerate(q.quarter, 2):
        c = w3.cell(1, j, qq); c.font = H; c.fill = FILL; w3.column_dimensions[col(j)].width = 9
        for i, m in enumerate(raw, 2):
            cell = w3.cell(i, j, float(q[m].iloc[j - 2])); cell.font = BLUE; cell.number_format = "0.00" if "pct" in m else "#,##0"
    for i, m in enumerate(raw, 2): w3.cell(i, 1, m).font = N
    for k, m in enumerate(["nii", "net_profit"]):
        r_src = raw.index(m) + 2
        w3.cell(9 + 2 * k, 1, f"{m} YoY").font = N; w3.cell(10 + 2 * k, 1, f"{m} QoQ").font = N
        for j in range(2, len(q) + 2):
            if j >= 6: w3.cell(9 + 2 * k, j, f"={col(j)}{r_src}/{col(j - 4)}{r_src}-1").number_format = "0.0%"
            if j >= 3: w3.cell(10 + 2 * k, j, f"={col(j)}{r_src}/{col(j - 1)}{r_src}-1").number_format = "0.0%"

    w4 = wb.create_sheet("Q2FY27_Results_Day")
    hdr = ["Metric (Rs crore, Screener basis)", "FY27E (model)", "Q1 FY27 actual", "Q2 FY27E = (FY27E - Q1) / 3", "Q2 FY26 actual", "Implied YoY",
           "Q2 FY27 ACTUAL (type here)", "Variance vs estimate", "Verdict (+/-3%)"]
    for j, t in enumerate(hdr, 1):
        c = w4.cell(1, j, t); c.font = H; c.fill = FILL; c.alignment = Alignment(wrap_text=True); w4.column_dimensions[col(j)].width = 16
    w4.column_dimensions["A"].width = 30
    qmap = {"NII (Screener basis)": "nii", "Other income": "other_income", "PBT": "pbt", "Net profit": "net_profit"}
    last, yago = q.iloc[-1], q[q.quarter == "Sep-25"].iloc[0]
    for i, r in enumerate(est.itertuples(index=False), 2):
        m = qmap[r.metric]
        vals = [r.metric, r.fy27e_model, float(last[m]), f"=(B{i}-C{i})/3", float(yago[m]), f"=D{i}/E{i}-1", None,
                f'=IF(G{i}="","",G{i}/D{i}-1)', f'=IF(G{i}="","",IF(H{i}>0.03,"Beat",IF(H{i}<-0.03,"Miss","In line")))']
        for j, v in enumerate(vals, 1):
            c = w4.cell(i, j, v); c.font = BLUE if j in (2, 3, 5) else N
            c.number_format = "0.0%" if j in (6, 8) else "#,##0"
        w4.cell(i, 7).fill = YEL
    w4.cell(7, 1, "Example of a filled row (Q1 FY27 net profit vs a naive Q4 FY26 run-rate):").font = Font(name="Arial", size=9, italic=True)
    for j, v in enumerate(["Net profit (example)", None, None, 725, None, None, 756, "=G8/D8-1", '=IF(H8>0.03,"Beat",IF(H8<-0.03,"Miss","In line"))'], 1):
        c = w4.cell(8, j, v); c.font = N; c.number_format = "0.0%" if j == 8 else "#,##0"
    w4.cell(10, 1, "Legend: blue = data, black = formula, yellow = type the reported number. Estimates spread the model's FY27E remainder evenly over Q2-Q4.").font = N
    wb.save(path)


def main():
    os.makedirs(f"{OUT}/charts", exist_ok=True)
    peers = pd.read_csv(f"{DATA}/peers.csv"); q = pd.read_csv(f"{DATA}/kvb_quarterly.csv"); est = pd.read_csv(f"{DATA}/kvb_estimates.csv")
    df, fit = build_comps(peers)
    chart_pe_growth(df, fit, f"{OUT}/charts/pe_vs_growth.png")
    pbfit, n_pb = pb_roe(df, f"{OUT}/charts/pb_vs_roe.png")
    q = build_tracker(q); chart_tracker(q, f"{OUT}/charts/kvb_quarterly.png")
    write_excel(df, fit, q, est, f"{OUT}/BFSI_Comps_Tracker.xlsx")

    k = df[df.bank == FOCUS].iloc[0]; med = df.groupby("size")["pe_ttm"].median(); last = q.iloc[-1]
    q2 = {r.metric: (r.fy27e_model - last[{"NII (Screener basis)": "nii", "Other income": "other_income", "PBT": "pbt", "Net profit": "net_profit"}[r.metric]]) / 3
          for r in est.itertuples()}
    lines = ["# BFSI comps & KVB tracker - summary", "",
             f"- Banks covered: {len(df)} (Large {sum(df['size'] == 'Large')}, Mid {sum(df['size'] == 'Mid')}, Small {sum(df['size'] == 'Small')}).",
             f"- Median P/E: all {df.pe_ttm.median():.1f}x | Large {med.get('Large', np.nan):.1f}x | Mid {med.get('Mid', np.nan):.1f}x | Small {med.get('Small', np.nan):.1f}x.",
             f"- {FOCUS}: P/E {k.pe_ttm:.1f}x ({k.pe_pctile:.0f}th percentile), Q1 FY27 profit growth {k.pat_yoy_pct:.1f}% ({k.growth_pctile:.0f}th percentile).",
             f"- P/E vs growth fit on {fit['n']} banks: slope {fit['b']:.3f}, R-squared {fit['r2']:.2f}. {FOCUS} residual {k.pe_residual:+.1f}x "
             f"({'cheaper' if k.pe_residual < 0 else 'richer'} than its growth implies).",
             "- Outliers excluded: " + ", ".join(df.loc[~df.in_fit, "bank"]) + " (turnarounds / base effects).",
             (f"- P/B vs ROE fit on {pbfit['n']} banks: R-squared {pbfit['r2']:.2f}." if pbfit else
              f"- P/B vs ROE regression inactive: only {n_pb} banks have book value and ROE. Fill bvps and roe_pct in data/peers.csv (Screener company pages) for 5+ banks."),
             f"- KVB {last.quarter}: NII {last.nii:,.0f} cr ({last.nii_yoy:+.0%} YoY), net profit {last.net_profit:,.0f} cr ({last.net_profit_yoy:+.0%} YoY), GNPA {last.gnpa_pct:.2f}%.",
             "- Q2 FY27 model-implied estimates (results 15-Oct-2026): " + ", ".join(f"{m} {v:,.0f} cr" for m, v in q2.items()) + "."]
    open(f"{OUT}/summary.md", "w").write("\n".join(lines) + "\n"); print("\n".join(lines))


if __name__ == "__main__":
    main()
